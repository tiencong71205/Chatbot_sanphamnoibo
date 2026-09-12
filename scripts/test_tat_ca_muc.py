"""
Test chatbot với TOÀN BỘ mục con của mọi sản phẩm.

Cách hoạt động: đọc trực tiếp file .docx trong data/source_docs, tự động
trích tên từng mục con (4.x = tính năng trên thiết bị, 5.x = chức năng với
ứng dụng, và các mục hướng dẫn lắp đặt), rồi sinh một câu hỏi cho mỗi mục
và gửi lên chatbot.

Vì danh sách câu hỏi lấy thẳng từ tài liệu chứ không viết cứng trong code,
tài liệu cập nhật thì bộ câu hỏi tự cập nhật theo — không bao giờ lệch.

Cách chạy:
    python3 scripts/test_tat_ca_muc.py                    # tất cả sản phẩm
    python3 scripts/test_tat_ca_muc.py "cảm biến hiện diện"   # lọc 1 sản phẩm
    python3 scripts/test_tat_ca_muc.py --liet-ke          # chỉ in câu hỏi, KHÔNG gọi chatbot

Kết quả in ra màn hình, đồng thời lưu vào ket_qua_test_muc.txt để đọc lại.
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path
from urllib.error import HTTPError, URLError

from docx import Document

BASE_URL = "http://localhost:8000"
SOURCE_DIR = Path(__file__).resolve().parent.parent / "data" / "source_docs"
OUTPUT_FILE = Path("ket_qua_test_muc.txt")
REQUEST_TIMEOUT = 120

# Mục con dạng "4.1.", "5.10.", "1.3." — dấu chấm giữa hai số.
SUBSECTION_RE = re.compile(r"^\s*(\d+)\.(\d+)\.?\s+(.+?)\s*$")
# Chương dạng "4. Tính năng trên thiết bị"
CHAPTER_RE = re.compile(r"^\s*(\d+)\.(?!\d)\s+(.+?)\s*$")

INSTALL_CHAPTER_RE = re.compile(r"hướng dẫn lắp đặt", re.IGNORECASE)
DEVICE_CHAPTER_RE = re.compile(r"tính năng trên thiết bị", re.IGNORECASE)
# Chấp nhận CẢ "Chức năng với ứng dụng" LẪN "Tính năng với ứng dụng":
# tài liệu không nhất quán -- hầu hết dùng "Chức năng", riêng file khóa
# cửa gỗ dùng "Tính năng". Mẫu cũ chỉ khớp "Chức năng" nên toàn bộ 15 mục
# 5.1-5.15 của sản phẩm đó chưa từng được test lần nào.
#
# Phải loại trừ "trên thiết bị" để không nuốt nhầm chương 4 ("Tính năng
# TRÊN THIẾT BỊ") -- hai chương này chỉ khác nhau ở mấy chữ cuối.
APP_CHAPTER_RE = re.compile(
    r"(chức năng|tính năng)(?!.*trên thiết bị).*ứng dụng", re.IGNORECASE
)


def _strip_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn"
    ).lower()


def is_heading(paragraph) -> bool:
    """Nhận diện tiêu đề: style Heading chuẩn, hoặc đoạn in đậm ngắn.

    Hai lớp vì tài liệu không nhất quán — một số mục dùng style Heading của
    Word, số khác chỉ là đoạn "List Paragraph" bôi đậm.
    """
    text = paragraph.text.strip()
    if not text or len(text) > 120:
        return False
    style = (paragraph.style.name or "") if paragraph.style else ""
    if style.startswith("Heading"):
        return True
    runs_bold = [r.bold for r in paragraph.runs if r.text.strip()]
    return bool(runs_bold) and all(runs_bold)


def extract_product_name(doc: Document) -> str:
    found_marker = False
    for para in doc.paragraphs:
        text = para.text.strip()
        if not found_marker and "TÀI LIỆU MÔ TẢ SẢN PHẨM" in text.upper():
            found_marker = True
            continue
        if found_marker and text and text.upper() == text and len(text) > 3:
            return text
    return "(không rõ tên sản phẩm)"


def extract_sections(path: Path) -> tuple[str, list[dict]]:
    """Trả về (tên sản phẩm, danh sách mục con cần hỏi).

    Mỗi mục con: {"so", "ten", "nhom"} — nhóm dùng để biết nên đặt câu hỏi
    theo kiểu nào (thao tác trên thiết bị / trên ứng dụng / lắp đặt).
    """
    doc = Document(str(path))
    product_name = extract_product_name(doc)

    sections: list[dict] = []
    current_group: str | None = None

    for para in doc.paragraphs:
        if not is_heading(para):
            continue
        text = para.text.strip()

        chapter = CHAPTER_RE.match(text)
        if chapter:
            title = chapter.group(2)
            if DEVICE_CHAPTER_RE.search(title):
                current_group = "thiet_bi"
            elif APP_CHAPTER_RE.search(title):
                current_group = "ung_dung"
            elif INSTALL_CHAPTER_RE.search(title):
                current_group = "lap_dat"
            else:
                current_group = None
            continue

        sub = SUBSECTION_RE.match(text)
        if sub and current_group:
            sections.append({
                "so": f"{sub.group(1)}.{sub.group(2)}",
                "ten": sub.group(3).strip().rstrip("."),
                "nhom": current_group,
            })

    return product_name, sections


def build_question(product_name: str, section: dict) -> str:
    """Sinh câu hỏi tự nhiên cho một mục con.

    Đặt câu theo đúng loại nội dung: mục lắp đặt hỏi "hướng dẫn", mục chức
    năng hỏi "cách thực hiện" — hỏi sai kiểu dễ khiến chatbot trả lời lệch
    trọng tâm dù đã tìm đúng tài liệu.
    """
    ten = section["ten"]
    if section["nhom"] == "lap_dat":
        return f"Hướng dẫn {ten.lower()} của {product_name}?"
    return f"{ten} của {product_name} thực hiện như thế nào?"


def ask_chatbot(question: str) -> str:
    payload = json.dumps({"question": question}).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE_URL}/api/chat",
        data=payload,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=REQUEST_TIMEOUT) as resp:
        return json.loads(resp.read().decode("utf-8")).get("answer", "(không có answer)")


def main(argv: list[str]) -> None:
    liet_ke_thoi = "--liet-ke" in argv
    filters = [a for a in argv if not a.startswith("--")]

    if not SOURCE_DIR.exists():
        print(f"Không tìm thấy thư mục tài liệu: {SOURCE_DIR}")
        return

    files = sorted(SOURCE_DIR.glob("*.docx"), key=lambda p: p.name)
    if not files:
        print(f"Không có file .docx nào trong {SOURCE_DIR}")
        return

    lines: list[str] = []

    def out(text: str = "") -> None:
        print(text)
        lines.append(text)

    tong_cau = 0
    for path in files:
        try:
            product_name, sections = extract_sections(path)
        except Exception as exc:  # noqa: BLE001 — muốn thấy lỗi rõ, không bỏ qua âm thầm
            out(f"\n[LỖI] Không đọc được {path.name}: {exc}")
            continue

        if filters and not any(
            _strip_accents(f) in _strip_accents(product_name) for f in filters
        ):
            continue

        if not sections:
            out(f"\n[!] {product_name}: không trích được mục con nào ({path.name})")
            continue

        out("\n" + "=" * 78)
        out(f"SẢN PHẨM: {product_name}   ({len(sections)} mục)")
        out("=" * 78)

        for section in sections:
            question = build_question(product_name, section)
            tong_cau += 1
            out(f"\n[{section['so']}] {section['ten']}")
            out(f"    Hỏi: {question}")

            if liet_ke_thoi:
                continue

            out("    " + "-" * 70)
            try:
                answer = ask_chatbot(question)
                out(answer)
            except HTTPError as e:
                out(f"    [LỖI HTTP {e.code}] {e.reason}")
            except URLError as e:
                out(f"    [LỖI KẾT NỐI] {e.reason} — chatbot có đang chạy không?")
            except Exception as e:  # noqa: BLE001
                out(f"    [LỖI] {e}")

    out("\n" + "=" * 78)
    out(f"TỔNG: {tong_cau} câu hỏi" + (" (chỉ liệt kê, chưa gọi chatbot)" if liet_ke_thoi else ""))

    try:
        OUTPUT_FILE.write_text("\n".join(lines), encoding="utf-8")
        print(f"\nĐã lưu kết quả vào: {OUTPUT_FILE.resolve()}")
    except OSError as exc:
        print(f"\n[!] Không lưu được file kết quả: {exc}")


if __name__ == "__main__":
    main(sys.argv[1:])
